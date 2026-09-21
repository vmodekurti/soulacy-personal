#!/usr/bin/env python3
"""JSON → .xlsx.  usage: xlsx_write.py DATA.json OUT.xlsx   DATA = {"Sheet": [[header,...],[row,...],...]}"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from _deps import need  # noqa: E402

openpyxl = need("openpyxl", "openpyxl")
from openpyxl.styles import Font  # noqa: E402

src, out = sys.argv[1:3]
data = json.load(open(src))
wb = openpyxl.Workbook()
wb.remove(wb.active)
for name, rows in data.items():
    ws = wb.create_sheet(title=str(name)[:31])
    for row in rows:
        ws.append(row)
    if rows:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
wb.save(out)
print(f"wrote {out} with {len(data)} sheet(s)")
